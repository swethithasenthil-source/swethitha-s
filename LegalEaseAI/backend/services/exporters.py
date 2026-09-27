from __future__ import annotations

import base64
import io
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF
from PIL import Image


def sanitize_text(text: str) -> str:
    if not text:
        return ""

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\u2022": "-",
        "\u2026": "...",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _safe_filename(value: str, extension: str) -> str:
    value = sanitize_text(value) or "LegalEase_Document"
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value)
    value = value.strip("._")

    if not value:
        value = "LegalEase_Document"

    return f"{value[:80]}.{extension}"


def _decode_logo(logo_base64: str | None) -> bytes | None:
    if not logo_base64:
        return None

    try:
        payload = (
            logo_base64.split(",", 1)[1]
            if "," in logo_base64
            else logo_base64
        )

        return base64.b64decode(payload)

    except Exception:
        return None


def build_txt(text: str):
    safe_text = sanitize_text(text)

    return (
        safe_text.encode("utf-8"),
        _safe_filename("LegalEase_Document", "txt"),
    )


def _set_docx_font(document: Document):
    styles = document.styles

    for style_name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
    ):
        style = styles[style_name]
        style.font.name = "Times New Roman"

        if style_name == "Normal":
            style.font.size = Pt(12)
        else:
            style.font.size = Pt(14)


def build_docx(
    text: str,
    document_type: str,
    terms: list[str],
    logo_base64: str | None = None,
):
    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    _set_docx_font(document)

    # Optional logo
    logo_bytes = _decode_logo(logo_base64)

    if logo_bytes:
        try:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run()

            run.add_picture(
                io.BytesIO(logo_bytes),
                width=Inches(1.2),
            )

        except Exception:
            pass

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        sanitize_text(document_type).upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    # Main document
    for raw_block in sanitize_text(text).split("\n\n"):

        block = raw_block.strip()

        if not block:
            continue

        lines = block.split("\n")
        first_line = lines[0].strip()

        # Heading
        if first_line.isupper() and len(first_line) < 100:

            paragraph = document.add_paragraph()

            heading_run = paragraph.add_run(first_line)

            heading_run.bold = True
            heading_run.font.name = "Times New Roman"
            heading_run.font.size = Pt(13)

            for line in lines[1:]:

                if line.strip():

                    paragraph2 = document.add_paragraph(
                        line.strip()
                    )

                    paragraph2.paragraph_format.space_after = Pt(6)

        else:

            paragraph = document.add_paragraph(block)

            paragraph.paragraph_format.space_after = Pt(8)

    # Key terms table
    clean_terms = [
        sanitize_text(term)
        for term in terms
        if sanitize_text(term)
    ]

    if clean_terms:

        document.add_paragraph()

        heading = document.add_paragraph()

        heading_run = heading.add_run("KEY TERMS")

        heading_run.bold = True

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, term in enumerate(
            clean_terms,
            start=1,
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = term

    # Footer
    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer.add_run(
        "Generated with LegalEase - Draft for review"
    )

    buffer = io.BytesIO()

    document.save(buffer)

    return (
        buffer.getvalue(),
        _safe_filename(document_type, "docx"),
    )


class LegalEasePDF(FPDF):

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=8,
        )

        self.cell(
            0,
            10,
            "Generated with LegalEase - Draft for review",
            align="C",
        )


def build_pdf(
    text: str,
    document_type: str,
    terms: list[str],
    logo_base64: str | None = None,
):

    pdf = LegalEasePDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=18,
    )

    pdf.set_margins(
        18,
        18,
        18,
    )

    pdf.add_page()

    # Optional logo
    logo_bytes = _decode_logo(logo_base64)

    temp_path = None

    if logo_bytes:

        try:

            image = Image.open(
                io.BytesIO(logo_bytes)
            )

            with tempfile.NamedTemporaryFile(
                delete=False,
                suffix=".png",
            ) as temp:

                image.save(
                    temp.name,
                    format="PNG",
                )

                temp_path = temp.name

            pdf.image(
                temp_path,
                x=85,
                y=18,
                w=40,
            )

            pdf.ln(45)

        except Exception:
            pass

    # Title
    pdf.set_font(
        "Helvetica",
        "B",
        16,
    )

    pdf.multi_cell(
        0,
        10,
        sanitize_text(document_type).upper(),
        align="C",
    )

    pdf.ln(5)

    # Document content
    for raw_block in sanitize_text(text).split("\n\n"):

        block = raw_block.strip()

        if not block:
            continue

        lines = block.split("\n")
        first_line = lines[0].strip()

        if first_line.isupper() and len(first_line) < 100:

            pdf.set_font(
                "Helvetica",
                "B",
                12,
            )

            pdf.multi_cell(
                0,
                7,
                first_line,
            )

            if len(lines) > 1:

                pdf.set_font(
                    "Helvetica",
                    size=11,
                )

                pdf.multi_cell(
                    0,
                    6,
                    "\n".join(lines[1:]).strip(),
                )

        else:

            pdf.set_font(
                "Helvetica",
                size=11,
            )

            pdf.multi_cell(
                0,
                6,
                block,
            )

        pdf.ln(2)

    # Key terms
    clean_terms = [
        sanitize_text(term)
        for term in terms
        if sanitize_text(term)
    ]

    if clean_terms:

        pdf.set_font(
            "Helvetica",
            "B",
            12,
        )

        pdf.multi_cell(
            0,
            7,
            "KEY TERMS",
        )

        pdf.set_font(
            "Helvetica",
            size=10,
        )

        for index, term in enumerate(
            clean_terms,
            start=1,
        ):

            pdf.multi_cell(
                0,
                6,
                f"{index}. {term}",
            )

    data = bytes(
        pdf.output(dest="S")
    )

    if temp_path:

        Path(temp_path).unlink(
            missing_ok=True
        )

    return (
        data,
        _safe_filename(document_type, "pdf"),
    )